from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.responses import StreamingResponse, FileResponse
from fastapi.middleware.cors import CORSMiddleware
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime, time
import io
import os

app = FastAPI(title="Sistema de Control Logístico de Transportes")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# 1. Ruta para el Procesador de Excel Principal
@app.get("/")
async def serve_index():
    return FileResponse("index.html")

# 2. Ruta para el Panel Administrativo de Reportes
@app.get("/reportes")
@app.get("/reportes.html")
async def serve_reportes():
    return FileResponse("reportes.html")

# 3. Ruta para el Formulario Móvil QR
@app.get("/evaluacion")
@app.get("/evaluacion.html")
async def serve_evaluacion():
    return FileResponse("evaluacion.html")
# 4. Ruta para generar el QR
@app.get("/qr_generator")
@app.get("/qr_generator.html")
async def serve_qr_generator():
    return FileResponse("qr_generator.html")

def parse_hora(hora_val):
    if pd.isna(hora_val):
        return None
    if isinstance(hora_val, time):
        return hora_val
    if isinstance(hora_val, datetime):
        return hora_val.time()
    
    hora_str = str(hora_val).strip()
    for fmt in ("%H:%M:%S", "%H:%M", "%I:%M:%S %p", "%I:%M %p"):
        try:
            return datetime.strptime(hora_str, fmt).time()
        except ValueError:
            pass
    return None

def clasificar_turno_por_hora(hora_val):
    h = parse_hora(hora_val)
    if h is None:
        return "ENTRADA TURNO GENERAL"

    sec = h.hour * 3600 + h.minute * 60 + h.second

    def to_sec(h_str):
        t = datetime.strptime(h_str, "%H:%M:%S").time()
        return t.hour * 3600 + t.minute * 60 + t.second

    if sec <= to_sec("05:59:59"):
        return "ENTRADA PRIMER TURNO"
    elif sec <= to_sec("06:29:59"):
        return "SALIDA TERCER TURNO"
    elif sec <= to_sec("12:29:59"):
        return "ENTRADA TURNO GENERAL"
    elif sec <= to_sec("13:59:59"):
        return "ENTRADA SEGUNDO TURNO"
    elif sec <= to_sec("16:29:59"):
        return "SALIDA PRIMER TURNO"
    elif sec <= to_sec("17:29:59"):
        return "ENTRADA TURNO LARGO"
    elif sec <= to_sec("18:29:59"):
        return "SALIDA TURNO LARGO"
    elif sec <= to_sec("20:19:59"):
        return "SALIDA TURNO GENERAL"
    elif sec <= to_sec("21:59:59"):
        return "ENTRADA TERCER TURNO"
    else:
        return "SALIDA SEGUNDO TURNO"

def limpiar_valor(val):
    if pd.isna(val):
        return ""
    if isinstance(val, float):
        if val.is_integer():
            return str(int(val))
        return str(val)
    return str(val)

@app.post("/procesar-reporte/")
async def procesar_reporte(file: UploadFile = File(...)):
    contents = await file.read()
    
    # Lectura inteligente de archivos (.xls, .xlsx, TSV/CSV)
    df = None
    try:
        df = pd.read_excel(io.BytesIO(contents))
    except Exception:
        try:
            df = pd.read_csv(io.BytesIO(contents), sep='\t')
        except Exception:
            try:
                df = pd.read_csv(io.BytesIO(contents))
            except Exception:
                raise HTTPException(status_code=400, detail="No se pudo interpretar la estructura del archivo cargado.")

    if df is None or df.empty:
        raise HTTPException(status_code=400, detail="El archivo cargado está vacío o no contiene datos válidos.")

    df.columns = [str(col).strip() for col in df.columns]

    # Priorización de la columna "Fecha" exacta
    col_fecha = "Fecha" if "Fecha" in df.columns else next((c for c in df.columns if str(c).lower() == 'fecha'), None)
    if not col_fecha:
        col_fecha = next((c for c in df.columns if 'fecha' in str(c).lower() and 'inicio' not in str(c).lower() and 'fin' not in str(c).lower()), None)

    col_hora = next((c for c in df.columns if 'hora' in str(c).lower()), None)
    col_camion = next((c for c in df.columns if 'clavecamion' in str(c).lower() or 'camion' in str(c).lower() or 'unidad' in str(c).lower()), None)
    col_cc = next((c for c in df.columns if 'centrodecostos' in str(c).lower() or 'centro_costos' in str(c).lower()), None)

    if not col_fecha or not col_hora or not col_camion:
        raise HTTPException(
            status_code=400, 
            detail="No se encontraron las columnas requeridas (Fecha, Hora, ClaveCamion)."
        )

    col_linea = next((c for c in df.columns if 'linea' in str(c).lower() and 'clave' not in str(c).lower() and 'id' not in str(c).lower()), None)
    if not col_linea:
        col_linea = next((c for c in df.columns if 'linea' in str(c).lower()), None)

    col_poblacion = next((c for c in df.columns if 'poblaci' in str(c).lower()), None)

    # Identificación de Bionatur
    if col_cc:
        df['Es_Bionatur'] = df[col_cc].astype(str).str.strip().str.endswith('D')
    else:
        df['Es_Bionatur'] = False

    df['Fecha_Clean'] = df[col_fecha].astype(str).str.split().str[0]
    df['Fecha_Parsed'] = pd.to_datetime(df['Fecha_Clean'], dayfirst=True, errors='coerce')
    
    if df['Fecha_Parsed'].dropna().empty:
        df['Fecha_Parsed'] = pd.to_datetime(df['Fecha_Clean'], errors='coerce')

    df['Turno_Clasificado'] = df[col_hora].apply(clasificar_turno_por_hora)
    df = df.sort_values(by=[col_camion, 'Fecha_Parsed', col_hora], ascending=[True, True, True])

    original_columns = [c for c in df.columns if c not in ['Es_Bionatur', 'Fecha_Clean', 'Fecha_Parsed', 'Turno_Clasificado']]

    wb = Workbook()
    wb.remove(wb.active)

    # --- MAPA DE COLORES DIFERENCIADOS POR TURNO ---
    turno_color_hex = {
        "ENTRADA PRIMER TURNO": "E3F2FD",   # Azul Suave
        "SALIDA TERCER TURNO": "F3E5F5",    # Morado Claro
        "ENTRADA TURNO GENERAL": "E8F5E9",  # Verde Menta
        "ENTRADA SEGUNDO TURNO": "FFF3E0",  # Naranja Claro
        "SALIDA PRIMER TURNO": "FFEBEE",   # Rosa / Coral
        "ENTRADA TURNO LARGO": "E0F7FA",   # Cian Suave
        "SALIDA TURNO LARGO": "FCE4EC",    # Magenta Suave
        "SALIDA TURNO GENERAL": "FFFDE7",  # Amarillo Pastel
        "ENTRADA TERCER TURNO": "E8EAF6",   # Índigo Suave
        "SALIDA SEGUNDO TURNO": "EFEBE9",   # Café Moka
    }

    turno_fills = {t: PatternFill(start_color=hex_val, end_color=hex_val, fill_type="solid") for t, hex_val in turno_color_hex.items()}

    # Estilos generales
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    summary_header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    bionatur_header_fill = PatternFill(start_color="1E4620", end_color="1E4620", fill_type="solid") # Verde pino profundo
    global_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    total_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    bionatur_total_fill = PatternFill(start_color="A5D6A7", end_color="A5D6A7", fill_type="solid")
    
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=11, bold=True)
    font_regular = Font(name="Calibri", size=11)
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')

    turnos_orden = [
        "ENTRADA PRIMER TURNO", "SALIDA TERCER TURNO", "ENTRADA TURNO GENERAL",
        "ENTRADA SEGUNDO TURNO", "SALIDA PRIMER TURNO", "ENTRADA TURNO LARGO",
        "SALIDA TURNO LARGO", "SALIDA TURNO GENERAL", "ENTRADA TERCER TURNO", 
        "SALIDA SEGUNDO TURNO"
    ]

    # --- 1. HOJA INICIAL DE RESUMEN GLOBAL ---
    ws_global = wb.create_sheet(title="RESUMEN GLOBAL", index=0)
    ws_global.views.sheetView[0].showGridLines = True

    ws_global.merge_cells("A1:F1")
    title_cell = ws_global["A1"]
    title_cell.value = "CONSOLIDADO LOGÍSTICO DE REPORTES PROCESADOS"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_cell.fill = global_fill
    title_cell.alignment = align_center

    global_headers = ["CAMIÓN / UNIDAD", "TOTAL REGISTROS", "TOTAL BIONATUR", "LÍNEA PREDOMINANTE", "POBLACIÓN", "DÍAS OPERATIVOS"]
    ws_global.append([])
    ws_global.append(global_headers)

    for col_idx in range(1, 7):
        c = ws_global.cell(row=3, column=col_idx)
        c.fill = summary_header_fill
        c.font = font_header
        c.alignment = align_center

    g_row = 4
    total_general_registros = 0
    total_general_bionatur = 0

    for camion_val, df_camion in df.groupby(col_camion, sort=True):
        camion_str = limpiar_valor(camion_val)
        cnt_total = len(df_camion)
        cnt_bionatur = int(df_camion['Es_Bionatur'].sum())

        total_general_registros += cnt_total
        total_general_bionatur += cnt_bionatur

        linea_str = ""
        if col_linea and not df_camion[col_linea].dropna().empty:
            linea_str = limpiar_valor(df_camion[col_linea].mode()[0])

        poblacion_str = ""
        if col_poblacion and not df_camion[col_poblacion].dropna().empty:
            poblacion_str = limpiar_valor(df_camion[col_poblacion].mode()[0])

        dias_ops = df_camion['Fecha_Parsed'].dropna().dt.date.nunique()

        ws_global.cell(row=g_row, column=1, value=camion_str).alignment = align_center
        ws_global.cell(row=g_row, column=2, value=cnt_total).alignment = align_right
        ws_global.cell(row=g_row, column=3, value=cnt_bionatur).alignment = align_right
        ws_global.cell(row=g_row, column=4, value=linea_str).alignment = align_left
        ws_global.cell(row=g_row, column=5, value=poblacion_str).alignment = align_left
        ws_global.cell(row=g_row, column=6, value=dias_ops).alignment = align_center

        for c_i in range(1, 7):
            cell = ws_global.cell(row=g_row, column=c_i)
            cell.font = font_regular
            cell.border = thin_border

        g_row += 1

    # Fila Total Global
    ws_global.cell(row=g_row, column=1, value="TOTAL GENERAL").alignment = align_center
    ws_global.cell(row=g_row, column=2, value=total_general_registros).alignment = align_right
    ws_global.cell(row=g_row, column=3, value=total_general_bionatur).alignment = align_right
    
    for c_i in range(1, 7):
        cell = ws_global.cell(row=g_row, column=c_i)
        cell.fill = total_fill
        cell.font = font_bold
        cell.border = thin_border

    # --- 2. GENERACIÓN DE HOJAS POR CAMIÓN ---
    for camion_val, df_camion in df.groupby(col_camion, sort=True):
        camion_str = limpiar_valor(camion_val)
        ws = wb.create_sheet(title=f"CAMION {camion_str}")
        ws.views.sheetView[0].showGridLines = True
        
        ws.append(original_columns)
        for col_num in range(1, len(original_columns) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = font_header
            cell.alignment = align_center

        row_pointer = 3
        fechas_unicas = sorted(df_camion['Fecha_Parsed'].dropna().unique())
        summary_col_start = len(original_columns) + 2
        bionatur_col_start = summary_col_start + 3  # Separación de 1 columna

        for fecha_p in fechas_unicas:
            df_fecha = df_camion[df_camion['Fecha_Parsed'] == fecha_p]
            fecha_actual = fecha_p.date()
            
            start_block_row = row_pointer
            current_left_row = start_block_row

            # Rellenar registros tabla izquierda
            for _, row in df_fecha.iterrows():
                es_bio = row['Es_Bionatur']
                turno_nombre = row['Turno_Clasificado']
                
                for col_idx, col_name in enumerate(original_columns, 1):
                    val = row[col_name]
                    if col_name == col_fecha and pd.notna(val):
                        val_str = fecha_actual.strftime('%Y-%m-%d')
                    else:
                        val_str = limpiar_valor(val)

                    cell = ws.cell(row=current_left_row, column=col_idx, value=val_str)
                    cell.font = font_regular
                    cell.border = thin_border
                    cell.alignment = align_center if col_name in [col_fecha, col_hora, col_camion] else align_left
                    
                    # Si es Bionatur, aplicar el color específico de su turno
                    if es_bio and turno_nombre in turno_fills:
                        cell.fill = turno_fills[turno_nombre]

                current_left_row += 1

            linea_nombre = ""
            if col_linea and not df_fecha[col_linea].dropna().empty:
                linea_nombre = limpiar_valor(df_fecha[col_linea].dropna().iloc[0])

            poblacion_valor = ""
            if col_poblacion and not df_fecha[col_poblacion].dropna().empty:
                poblacion_valor = limpiar_valor(df_fecha[col_poblacion].dropna().iloc[0])

            conteo_turnos_general = df_fecha['Turno_Clasificado'].value_counts()
            
            df_fecha_bionatur = df_fecha[df_fecha['Es_Bionatur'] == True]
            conteo_turnos_bionatur = df_fecha_bionatur['Turno_Clasificado'].value_counts()

            # Encabezados Tabla General
            c_hdr1 = ws.cell(row=start_block_row, column=summary_col_start, value="CONCEPTO / TURNO GENERAL")
            v_hdr1 = ws.cell(row=start_block_row, column=summary_col_start + 1, value="CONTEO")
            for h_cell in [c_hdr1, v_hdr1]:
                h_cell.fill = summary_header_fill
                h_cell.font = font_header
                h_cell.alignment = align_center

            # Encabezados Tabla Bionatur
            c_hdr2 = ws.cell(row=start_block_row, column=bionatur_col_start, value="CONCEPTO / TURNO BIONATUR")
            v_hdr2 = ws.cell(row=start_block_row, column=bionatur_col_start + 1, value="CONTEO")
            for h_cell in [c_hdr2, v_hdr2]:
                h_cell.fill = bionatur_header_fill
                h_cell.font = font_header
                h_cell.alignment = align_center

            summary_items_gen = [
                ("LINEA", linea_nombre),
                ("FECHA", fecha_actual.strftime('%Y-%m-%d')),
                ("POBLACIÓN", poblacion_valor),
                ("CAMIÓN", camion_str)
            ]
            summary_items_bio = [
                ("LINEA", linea_nombre),
                ("FECHA", fecha_actual.strftime('%Y-%m-%d')),
                ("POBLACIÓN", poblacion_valor),
                ("CAMIÓN", camion_str)
            ]

            total_gen = 0
            total_bio = 0
            for t in turnos_orden:
                cnt_g = int(conteo_turnos_general.get(t, 0))
                cnt_b = int(conteo_turnos_bionatur.get(t, 0))
                
                summary_items_gen.append((t, cnt_g))
                summary_items_bio.append((t, cnt_b))
                
                total_gen += cnt_g
                total_bio += cnt_b

            summary_items_gen.append(("TOTAL REGISTROS", total_gen))
            summary_items_bio.append(("TOTAL BIONATUR", total_bio))

            # Rellenar Tabla 1 (General)
            s_row = start_block_row + 1
            for concept, val in summary_items_gen:
                c_cell = ws.cell(row=s_row, column=summary_col_start, value=concept)
                v_cell = ws.cell(row=s_row, column=summary_col_start + 1, value=val)

                c_cell.border = thin_border
                v_cell.border = thin_border

                if concept in turno_fills:
                    c_cell.fill = turno_fills[concept]
                    v_cell.fill = turno_fills[concept]

                if concept == "TOTAL REGISTROS":
                    c_cell.fill = total_fill
                    v_cell.fill = total_fill
                    c_cell.font = font_bold
                    v_cell.font = font_bold
                    v_cell.alignment = align_right
                else:
                    c_cell.font = font_bold if s_row <= start_block_row + 4 else font_regular
                    v_cell.font = font_regular
                    v_cell.alignment = align_left if isinstance(val, str) else align_right

                s_row += 1

            # Rellenar Tabla 2 (Bionatur) con colores únicos por turno
            s_row = start_block_row + 1
            for concept, val in summary_items_bio:
                c_cell = ws.cell(row=s_row, column=bionatur_col_start, value=concept)
                v_cell = ws.cell(row=s_row, column=bionatur_col_start + 1, value=val)

                c_cell.border = thin_border
                v_cell.border = thin_border

                # Aplicar color pastel específico del turno
                if concept in turno_fills:
                    c_cell.fill = turno_fills[concept]
                    v_cell.fill = turno_fills[concept]

                if concept == "TOTAL BIONATUR":
                    c_cell.fill = bionatur_total_fill
                    v_cell.fill = bionatur_total_fill
                    c_cell.font = font_bold
                    v_cell.font = font_bold
                    v_cell.alignment = align_right
                else:
                    c_cell.font = font_bold if s_row <= start_block_row + 4 else font_regular
                    v_cell.font = font_regular
                    v_cell.alignment = align_left if isinstance(val, str) else align_right

                s_row += 1

            max_rows = max(len(df_fecha), len(summary_items_gen) + 1)
            row_pointer = start_block_row + max_rows + 3

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    for col in ws_global.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws_global.column_dimensions[col_letter].width = max(max_len + 4, 15)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=Reporte_Procesado_{datetime.now().strftime('%Y%m%d')}.xlsx"}
    )

if __name__ == "__main__":
    import uvicorn
    port = int(os.environ.get("PORT", 8000))
    uvicorn.run(app, host="0.0.0.0", port=port)